from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "POLICY_CLASS_Q_MEMORY_MECHANISM_PAPER.md"
OUTPUT = ROOT / "POLICY_CLASS_Q_MEMORY_MECHANISM_PAPER.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=11, color="222222", bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text, size=11, color="222222"):
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]), size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color, italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, color=color, name="Courier New")
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size, color=color)


def set_paragraph(p, before=0, after=8, line=1.333, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    p.alignment = align


def add_body(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p)
    add_inline(p, text)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    if level == 1:
        set_paragraph(p, before=18, after=10, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
        size, color = 16, "2E74B5"
    elif level == 2:
        set_paragraph(p, before=12, after=6, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
        size, color = 13, "2E74B5"
    else:
        set_paragraph(p, before=8, after=4, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
        size, color = 12, "1F4D78"
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=True)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p, before=4, after=8, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    set_run_font(run, size=10, color="333333", name="Cambria Math")
    return p


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    widths = [3800] + [int((9360 - 3800) / (len(rows[0]) - 1))] * (len(rows[0]) - 1)
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph(p, before=0, after=0, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            add_inline(p, value, size=9.5)
            if ri == 0:
                set_cell_shading(cell, "E8EEF5")
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_title_block(doc):
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run("Decision-Preserving Discrete Memory via Policy-Class Q-Distortion")
    set_run_font(run, size=22, color="0B2545", bold=True)
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=3, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run("A mechanism and methods study of rate-constrained memory for delayed decisions")
    set_run_font(run, size=12.5, color="555555", italic=True)
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=16, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run("Evidence-grounded manuscript draft | Open-ended agent experiments excluded by design")
    set_run_font(run, size=9.5, color="777777")


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("Policy-Class Q-Sufficient Memory | Mechanism paper"), size=8.5, color="777777")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("Evidence-grounded draft"), size=8.5, color="777777")


def build():
    lines = SOURCE.read_text().splitlines()
    doc = Document()
    configure_doc(doc)
    add_title_block(doc)
    in_abstract = False
    idx = 0
    while idx < len(lines):
        line = lines[idx].rstrip()
        if not line:
            idx += 1
            continue
        if line.startswith("# "):
            idx += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            in_abstract = title.lower() == "abstract"
            add_heading(doc, title, 1)
            idx += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
            idx += 1
            continue
        if line.startswith("#### "):
            add_heading(doc, line[5:].strip(), 3)
            idx += 1
            continue
        if line.startswith("|") and idx + 1 < len(lines) and lines[idx + 1].startswith("|"):
            rows = []
            while idx < len(lines) and lines[idx].startswith("|"):
                row = [part.strip() for part in lines[idx].strip().strip("|").split("|")]
                if not all(set(part) <= set("-:") for part in row):
                    rows.append(row)
                idx += 1
            add_table(doc, rows)
            continue
        if line.startswith("\\["):
            equation = line[2:].strip()
            idx += 1
            while idx < len(lines) and not lines[idx].strip().endswith("\\]"):
                equation += " " + lines[idx].strip()
                idx += 1
            if idx < len(lines):
                equation += " " + lines[idx].strip()[:-2].strip()
            add_equation(doc, equation)
            idx += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            set_paragraph(p, before=4, after=8, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.left_indent = Inches(0.25)
            add_inline(p, line[2:].strip(), size=10.5, color="1F3A5F")
            idx += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph(p, before=0, after=4, line=1.208, align=WD_ALIGN_PARAGRAPH.LEFT)
            add_inline(p, line[2:].strip())
            idx += 1
            continue
        add_body(doc, line)
        idx += 1
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
